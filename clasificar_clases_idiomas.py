from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import time
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


FORMATOS_AUDIO = {".aac", ".flac", ".m4a", ".mp3", ".mp4", ".mpeg", ".mpga", ".ogg", ".wav", ".webm"}
FORMATOS_TEXTO = {".txt", ".md"}
MODELO_AUDIO = "whisper-large-v3-turbo"
MODELO_TEXTO = "openai/gpt-oss-20b"
VERSION_ANALISIS = 3

CODIGOS_ISO_639_1 = frozenset(
    """aa ab ae af ak am an ar as av ay az ba be bg bh bi bm bn bo br bs
    ca ce ch co cr cs cu cv cy da de dv dz ee el en eo es et eu fa ff fi fj
    fo fr fy ga gd gl gn gu gv ha he hi ho hr ht hu hy hz ia id ie ig ii ik
    io is it iu ja jv ka kg ki kj kk kl km kn ko kr ks ku kv kw ky la lb lg
    li ln lo lt lu lv mg mh mi mk ml mn mr ms mt my na nb nd ne ng nl nn no
    nr nv ny oc oj om or os pa pi pl ps pt qu rm rn ro ru rw sa sc sd se sg
    si sk sl sm sn so sq sr ss st su sv sw ta te tg th ti tk tl tn to tr ts
    tt tw ty ug uk ur uz ve vi vo wa wo xh yi yo za zh zu""".split()
)

ALIAS_CODIGOS_IDIOMA = {
    "gr": "el", "gre": "el", "ell": "el",
    "jp": "ja", "jpn": "ja",
    "sp": "es", "spa": "es",
    "eng": "en", "fre": "fr", "fra": "fr",
    "ger": "de", "deu": "de", "ita": "it", "tur": "tr",
    "por": "pt", "ara": "ar", "chi": "zh", "zho": "zh",
    "kor": "ko", "rus": "ru", "dut": "nl", "nld": "nl",
    "pol": "pl", "swe": "sv", "ukr": "uk", "rum": "ro", "ron": "ro",
    "cze": "cs", "ces": "cs", "dan": "da", "fin": "fi",
    "heb": "he", "hin": "hi", "ind": "id", "vie": "vi", "tha": "th",
    "cat": "ca", "baq": "eu", "eus": "eu",
    "in": "id", "iw": "he", "ji": "yi",
}

NOMBRES_IDIOMA = {
    "arabic": "ar",
    "arabe": "ar",
    "árabe": "ar",
    "catalan": "ca",
    "catalán": "ca",
    "checo": "cs",
    "czech": "cs",
    "danish": "da",
    "danés": "da",
    "deutsch": "de",
    "dutch": "nl",
    "english": "en",
    "espanol": "es",
    "español": "es",
    "francais": "fr",
    "français": "fr",
    "french": "fr",
    "german": "de",
    "ingles": "en",
    "inglés": "en",
    "italian": "it",
    "italiano": "it",
    "portugues": "pt",
    "português": "pt",
    "portuguese": "pt",
    "spanish": "es",
    "aleman": "de",
    "alemán": "de",
    "chino": "zh",
    "chinese": "zh",
    "frances": "fr",
    "francés": "fr",
    "greek": "el",
    "griego": "el",
    "hebreo": "he",
    "hebrew": "he",
    "hindi": "hi",
    "indonesian": "id",
    "indonesio": "id",
    "japanese": "ja",
    "japones": "ja",
    "japonés": "ja",
    "korean": "ko",
    "neerlandés": "nl",
    "noruego": "no",
    "norwegian": "no",
    "persa": "fa",
    "polaco": "pl",
    "polish": "pl",
    "rumano": "ro",
    "romanian": "ro",
    "ruso": "ru",
    "russian": "ru",
    "turco": "tr",
    "turkish": "tr",
    "türkçe": "tr",
    "sueco": "sv",
    "swedish": "sv",
    "tailandés": "th",
    "thai": "th",
    "ucraniano": "uk",
    "ukrainian": "uk",
    "vietnamese": "vi",
    "vietnamita": "vi",
}

IDIOMAS_EN_ESPANOL = {
    "ar": "árabe", "ca": "catalán", "cs": "checo", "da": "danés",
    "de": "alemán", "el": "griego", "en": "inglés", "es": "español",
    "fa": "persa", "fi": "finés", "fr": "francés", "he": "hebreo",
    "hi": "hindi", "id": "indonesio", "it": "italiano", "ja": "japonés",
    "ko": "coreano", "nl": "neerlandés", "no": "noruego", "pl": "polaco",
    "pt": "portugués", "ro": "rumano", "ru": "ruso", "sv": "sueco",
    "th": "tailandés", "tr": "turco", "uk": "ucraniano", "vi": "vietnamita",
    "zh": "chino",
    "desconocido": "sin identificar",
}

AZUL_OSCURO = "14283D"
TURQUESA = "19B5C5"
BLANCO = "FFFFFF"
TEXTO_OSCURO = "1F2933"

def leer_argumentos():
    parser = argparse.ArgumentParser(
        description="Lee o transcribe clases, identifica los idiomas y extrae los temas enseñados."
    )
    parser.add_argument("--entrada", default="entrada")
    parser.add_argument("--salida", default="salida")
    parser.add_argument(
        "--idioma",
        default="auto",
        help="Idioma ISO-639-1 (por ejemplo, 'es') o 'auto' para detectarlo. Predeterminado: auto.",
    )
    parser.add_argument("--reprocesar", action="store_true")
    parser.add_argument("--sin-copiar", action="store_true")
    return parser.parse_args()


def resolver(base: Path, nombre: str) -> Path:
    ruta = Path(nombre)
    return ruta.resolve() if ruta.is_absolute() else (base / ruta).resolve()


def leer_estado(ruta: Path):
    if ruta.exists():
        try:
            datos = json.loads(ruta.read_text(encoding="utf-8"))
            if isinstance(datos.get("procesados"), dict):
                return datos
        except (OSError, json.JSONDecodeError):
            pass
    return {"procesados": {}}


def guardar_json(ruta: Path, datos) -> None:
    temporal = ruta.with_suffix(ruta.suffix + ".tmp")
    temporal.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    temporal.replace(ruta)


def firma(ruta: Path) -> str:
    datos = ruta.stat()
    texto = f"{ruta.resolve()}|{datos.st_size}|{datos.st_mtime_ns}"
    return hashlib.sha256(texto.encode()).hexdigest()


def hash_contenido(ruta: Path) -> str:
    resumen = hashlib.sha256()
    with ruta.open("rb") as archivo:
        while bloque := archivo.read(1024 * 1024):
            resumen.update(bloque)
    return resumen.hexdigest()


def normalizar_idioma(valor) -> str:
    texto = str(valor or "").strip().lower().replace("_", "-")
    if texto in NOMBRES_IDIOMA:
        return NOMBRES_IDIOMA[texto]
    if texto in ALIAS_CODIGOS_IDIOMA:
        return ALIAS_CODIGOS_IDIOMA[texto]
    codigo = texto.split("-", 1)[0]
    if codigo in ALIAS_CODIGOS_IDIOMA:
        return ALIAS_CODIGOS_IDIOMA[codigo]
    if codigo in CODIGOS_ISO_639_1:
        return codigo
    return "desconocido"


def segundos_espera_groq(error: Exception, intento: int) -> float:
    respuesta = getattr(error, "response", None)
    encabezados = getattr(respuesta, "headers", None)
    if encabezados:
        valor = encabezados.get("retry-after")
        try:
            return max(1.0, min(60.0, float(valor) + 0.5))
        except (TypeError, ValueError):
            pass

    mensaje = str(error)
    coincidencia = re.search(
        r"try again in\s+(?:(\d+)m)?\s*([0-9.]+)s",
        mensaje,
        re.IGNORECASE,
    )
    if coincidencia:
        minutos = int(coincidencia.group(1) or 0)
        segundos = float(coincidencia.group(2))
        return max(1.0, min(60.0, minutos * 60 + segundos + 0.5))
    coincidencia_ms = re.search(r"try again in\s+([0-9.]+)ms", mensaje, re.IGNORECASE)
    if coincidencia_ms:
        return max(1.0, min(60.0, float(coincidencia_ms.group(1)) / 1000 + 0.5))
    return float(min(60, 3 * intento))


def llamar_groq_con_reintentos(accion, descripcion: str, reintentar_json: bool = False):
    for intento in range(1, 5):
        try:
            return accion()
        except Exception as error:
            mensaje = str(error).lower()
            codigo = getattr(error, "status_code", None)
            reintentable = codigo in {408, 409, 429, 500, 502, 503, 504} or any(
                texto in mensaje
                for texto in ("rate limit", "timeout", "timed out", "connection error")
            )
            if reintentar_json and "json_validate_failed" in mensaje:
                reintentable = True
            if not reintentable or intento == 4:
                raise
            espera = segundos_espera_groq(error, intento)
            print(
                f"  Groq pidió esperar durante {descripcion}; "
                f"reintento {intento}/3 en {espera:.1f} s..."
            )
            time.sleep(espera)


def transcribir(cliente: Groq, audio: Path, idioma: str) -> tuple[str, str]:
    opciones = {"model": MODELO_AUDIO, "response_format": "verbose_json", "temperature": 0}
    if idioma.lower() != "auto":
        idioma_normalizado = normalizar_idioma(idioma)
        if idioma_normalizado == "desconocido":
            raise ValueError(f"Código de idioma no válido para Whisper: {idioma!r}")
        opciones["language"] = idioma_normalizado
    contenido = audio.read_bytes()
    respuesta = llamar_groq_con_reintentos(
        lambda: cliente.audio.transcriptions.create(
            file=(audio.name, contenido),
            **opciones,
        ),
        f"la transcripción de {audio.name}",
    )
    texto = getattr(respuesta, "text", "")
    if not texto or not texto.strip():
        raise RuntimeError("La transcripción quedó vacía.")
    idioma_detectado = idioma if idioma.lower() != "auto" else getattr(respuesta, "language", "")
    return texto.strip(), normalizar_idioma(idioma_detectado)


def leer_texto(ruta: Path) -> str:
    ultimo_error = None
    for codificacion in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            texto = ruta.read_text(encoding=codificacion)
            if not texto.strip():
                raise RuntimeError("El archivo de texto está vacío.")
            return texto.strip()
        except UnicodeDecodeError as error:
            ultimo_error = error
    raise RuntimeError(f"No pude leer la codificación del texto: {ultimo_error}")


def destino_libre(ruta: Path) -> Path:
    if not ruta.exists():
        return ruta
    numero = 2
    while True:
        candidato = ruta.with_name(f"{ruta.stem}_{numero}{ruta.suffix}")
        if not candidato.exists():
            return candidato
        numero += 1


def normalizar_lista_idiomas(valor) -> list[str]:
    if isinstance(valor, list):
        candidatos = valor
    elif valor:
        candidatos = re.split(r"[,;|]", str(valor))
    else:
        candidatos = []
    salida = []
    for candidato in candidatos:
        codigo = normalizar_idioma(candidato)
        if codigo != "desconocido" and codigo not in salida:
            salida.append(codigo)
    return salida


def nombre_idioma(codigo: str) -> str:
    return IDIOMAS_EN_ESPANOL.get(codigo, codigo)


def nombre_carpeta(texto: str) -> str:
    sin_tildes = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    limpio = re.sub(r"[^a-zA-Z0-9_-]+", "_", sin_tildes.strip().lower()).strip("_")
    return limpio or "sin_identificar"


def identificador_archivo(ruta: Path) -> str:
    return hashlib.sha256(str(ruta.resolve()).encode()).hexdigest()[:12]


def abreviar_clase(texto: str, maximo: int = 24_000) -> tuple[str, bool]:
    if len(texto) <= maximo:
        return texto, False
    separador = "\n\n[... parte omitida por longitud ...]\n\n"
    disponible = maximo - 2 * len(separador)
    parte = disponible // 3
    parte_central = disponible - 2 * parte
    centro = len(texto) // 2
    mitad = parte_central // 2
    abreviado = (
        texto[:parte]
        + separador
        + texto[centro - mitad:centro - mitad + parte_central]
        + separador
        + texto[-parte:]
    )
    return abreviado, True


def analizar_clase(
    cliente: Groq,
    texto: str,
    idioma_sugerido: str,
    nombre_archivo: str = "",
) -> dict:
    texto_analizado, texto_recortado = abreviar_clase(texto)
    if texto_recortado:
        origen = f" de {nombre_archivo}" if nombre_archivo else ""
        print(
            f"  AVISO: la transcripción{origen} tiene {len(texto):,} caracteres. "
            f"El modelo analizará {len(texto_analizado):,}; el TXT completo se conserva."
        )
    instruccion = f"""Analiza esta transcripción de una clase de idiomas.

Distingue cuidadosamente:
1. idiomas_hablados: idiomas usados para explicar o presentar la clase.
2. idiomas_ensenados: idiomas que el estudiante está aprendiendo, aunque la explicación sea en otro idioma.
3. temas: uno o varios temas concretos enseñados. No uses etiquetas vagas como "gramática"
   si puedes indicar "verbo estar", "adjetivos posesivos", "presentaciones y saludos"
   o "formación de preguntas". Escribe los temas en español.

Cada elemento de temas debe ser un objeto con idioma (código ISO-639-1 del idioma enseñado)
y tema. Si hay varios idiomas o varios temas, incluye todos los que tengan evidencia.
Acepta cualquier idioma del mundo, no solo los ejemplos. Usa siempre su código ISO-639-1;
si aparece un idioma nuevo, inclúyelo normalmente para que Python cree su carpeta.
Idioma hablado sugerido por Whisper: {idioma_sugerido}.

Devuelve solamente JSON válido con esta forma:
{{
  "idiomas_hablados": ["en"],
  "idiomas_ensenados": ["tr"],
  "temas": [
    {{"idioma": "tr", "tema": "verbo estar"}},
    {{"idioma": "tr", "tema": "adjetivos"}}
  ],
  "resumen": "Resumen breve en español"
}}

TRANSCRIPCIÓN:
{texto_analizado}"""
    opciones = {
        "model": MODELO_TEXTO,
        "messages": [
            {
                "role": "system",
                "content": "Eres un analista de clases de idiomas. Responde solo con JSON válido.",
            },
            {"role": "user", "content": instruccion},
        ],
        "response_format": {"type": "json_object"},
        "temperature": 0,
    }
    respuesta = llamar_groq_con_reintentos(
        lambda: cliente.chat.completions.create(**opciones),
        "el análisis de la clase",
        reintentar_json=True,
    )
    resultado = json.loads(respuesta.choices[0].message.content or "{}")
    hablados = normalizar_lista_idiomas(resultado.get("idiomas_hablados"))
    if idioma_sugerido != "desconocido" and idioma_sugerido not in hablados:
        hablados.insert(0, idioma_sugerido)
    ensenados = normalizar_lista_idiomas(resultado.get("idiomas_ensenados"))

    temas = []
    vistos = set()
    temas_crudos = resultado.get("temas", [])
    if not isinstance(temas_crudos, list):
        temas_crudos = [temas_crudos]
    for elemento in temas_crudos[:20]:
        if isinstance(elemento, dict):
            idioma_tema = normalizar_idioma(elemento.get("idioma"))
            tema = str(elemento.get("tema") or elemento.get("nombre") or "").strip()
        else:
            idioma_tema = ensenados[0] if len(ensenados) == 1 else "desconocido"
            tema = str(elemento or "").strip()
        if idioma_tema == "desconocido" and len(ensenados) == 1:
            idioma_tema = ensenados[0]
        if idioma_tema != "desconocido" and idioma_tema not in ensenados:
            ensenados.append(idioma_tema)
        clave_tema = (idioma_tema, tema.casefold())
        if tema and clave_tema not in vistos:
            vistos.add(clave_tema)
            temas.append({"idioma": idioma_tema, "tema": tema})

    if not ensenados:
        ensenados = ["desconocido"]
    if not hablados:
        hablados = ["desconocido"]
    if not temas:
        temas = [{"idioma": ensenados[0], "tema": "tema no identificado"}]
    return {
        "idiomas_hablados": hablados,
        "idiomas_ensenados": ensenados,
        "temas": temas,
        "resumen": str(resultado.get("resumen", "")).strip(),
        "texto_recortado": texto_recortado,
        "caracteres_transcripcion": len(texto),
        "caracteres_analizados": len(texto_analizado),
    }


def copiar_por_idioma(archivo: Path, salida_clases: Path, idiomas: list[str], hash_actual: str) -> None:
    for codigo in idiomas:
        carpeta = salida_clases / nombre_carpeta(nombre_idioma(codigo))
        carpeta.mkdir(parents=True, exist_ok=True)
        destino = carpeta / archivo.name
        if destino.exists() and hash_contenido(destino) == hash_actual:
            continue
        shutil.copy2(archivo, destino_libre(destino))


def cargar_detalles_nuevos(ruta: Path) -> list[dict]:
    detalles = []
    if not ruta.exists():
        return detalles
    for archivo in ruta.glob("*.json"):
        try:
            dato = json.loads(archivo.read_text(encoding="utf-8"))
            if dato.get("version_analisis") == VERSION_ANALISIS:
                detalles.append(dato)
        except (OSError, json.JSONDecodeError):
            continue
    return detalles


def configurar_hoja(hoja, titulo: str) -> None:
    hoja.sheet_view.showGridLines = False
    hoja["A2"] = titulo
    hoja["A2"].font = Font(name="Arial", size=14, bold=True, color=AZUL_OSCURO)
    hoja["A3"].fill = PatternFill("solid", fgColor=AZUL_OSCURO)
    hoja["A3"].border = Border(bottom=Side(style="thin", color=TURQUESA))
    hoja.row_dimensions[2].height = 22


def estilo_encabezado(celdas) -> None:
    for celda in celdas:
        celda.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
        celda.font = Font(name="Arial", size=10, bold=True, color=BLANCO)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def agregar_tabla(hoja, referencia: str, nombre: str) -> None:
    tabla = Table(displayName=nombre, ref=referencia)
    tabla.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    hoja.add_table(tabla)


def crear_excel(ruta: Path, detalles: list[dict]) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    clases = libro.create_sheet("Clases")
    temas_hoja = libro.create_sheet("Temas por idioma")

    resumen_por_idioma = defaultdict(lambda: {"archivos": set(), "temas": set()})
    filas_clases = []
    filas_temas = []
    for detalle in sorted(
        detalles,
        key=lambda x: (x.get("procesado_utc", ""), x.get("archivo_original", "")),
    ):
        archivo = Path(detalle.get("archivo_original", "")).name
        hablados = [nombre_idioma(x) for x in detalle.get("idiomas_hablados", [])]
        ensenados = [nombre_idioma(x) for x in detalle.get("idiomas_ensenados", [])]
        temas = detalle.get("temas", [])
        nombres_temas = [str(x.get("tema", "")) for x in temas if isinstance(x, dict)]
        fecha = detalle.get("procesado_utc", "")
        try:
            fecha_excel = datetime.fromisoformat(fecha.replace("Z", "+00:00")).replace(tzinfo=None)
        except (TypeError, ValueError):
            fecha_excel = fecha
        if detalle.get("texto_recortado"):
            cobertura = (
                f"Recortado: {detalle.get('caracteres_analizados', 0):,} de "
                f"{detalle.get('caracteres_transcripcion', 0):,} caracteres"
            )
        else:
            cobertura = "Completo"
        filas_clases.append([
            archivo,
            detalle.get("tipo_entrada", ""),
            ", ".join(hablados),
            ", ".join(ensenados),
            "; ".join(nombres_temas),
            detalle.get("resumen", ""),
            cobertura,
            fecha_excel,
        ])
        for codigo in detalle.get("idiomas_ensenados", []):
            resumen_por_idioma[codigo]["archivos"].add(archivo)
        for item in temas:
            if not isinstance(item, dict):
                continue
            codigo = normalizar_idioma(item.get("idioma"))
            tema = str(item.get("tema", "")).strip()
            resumen_por_idioma[codigo]["archivos"].add(archivo)
            if tema:
                resumen_por_idioma[codigo]["temas"].add(tema)
            filas_temas.append([
                nombre_idioma(codigo),
                tema,
                archivo,
                ", ".join(hablados),
            ])

    filas_temas.sort(
        key=lambda fila: (
            str(fila[0]).casefold(),
            str(fila[1]).casefold(),
            str(fila[2]).casefold(),
        )
    )

    configurar_hoja(resumen, "Clases de idiomas")
    resumen["A3"] = "Consolidado por idioma enseñado"
    for columna, valor in enumerate(["Idioma enseñado", "Clases", "Temas identificados"], 1):
        resumen.cell(row=5, column=columna, value=valor)
    estilo_encabezado(resumen[5][0:3])
    filas_resumen = []
    for codigo, datos in sorted(resumen_por_idioma.items(), key=lambda x: nombre_idioma(x[0])):
        filas_resumen.append([
            nombre_idioma(codigo),
            len(datos["archivos"]),
            "; ".join(sorted(datos["temas"], key=str.casefold)),
        ])
    if not filas_resumen:
        filas_resumen = [["Sin datos", 0, ""]]
    for fila in filas_resumen:
        resumen.append(fila)
    fin_resumen = 5 + len(filas_resumen)
    agregar_tabla(resumen, f"A5:C{fin_resumen}", "ResumenIdiomas")
    resumen.column_dimensions["A"].width = 22
    resumen.column_dimensions["B"].width = 12
    resumen.column_dimensions["C"].width = 70
    resumen.freeze_panes = "A6"
    for fila in resumen.iter_rows(min_row=6, max_row=fin_resumen, min_col=1, max_col=3):
        lineas = max(2, (len(str(fila[2].value or "")) + 69) // 70)
        resumen.row_dimensions[fila[0].row].height = min(180, 16 * lineas)
        for celda in fila:
            celda.font = Font(name="Arial", size=10, color=TEXTO_OSCURO)
            celda.alignment = Alignment(vertical="top", wrap_text=celda.column == 3)
    if filas_resumen[0][0] != "Sin datos":
        grafico = BarChart()
        grafico.title = "Clases por idioma enseñado"
        grafico.y_axis.title = "Clases"
        grafico.x_axis.title = "Idioma"
        grafico.style = 10
        grafico.height = 7
        grafico.width = 13
        grafico.add_data(
            Reference(resumen, min_col=2, min_row=5, max_row=fin_resumen),
            titles_from_data=True,
        )
        grafico.series[0].graphicalProperties.solidFill = TURQUESA
        grafico.series[0].graphicalProperties.line.solidFill = TURQUESA
        grafico.set_categories(Reference(resumen, min_col=1, min_row=6, max_row=fin_resumen))
        grafico.legend = None
        resumen.add_chart(grafico, "E5")

    configurar_hoja(clases, "Detalle de clases")
    encabezados_clases = [
        "Archivo", "Tipo", "Idioma hablado", "Idioma enseñado", "Temas",
        "Resumen", "Cobertura del texto", "Procesado UTC",
    ]
    for columna, valor in enumerate(encabezados_clases, 1):
        clases.cell(row=4, column=columna, value=valor)
    estilo_encabezado(clases[4][0:len(encabezados_clases)])
    for fila in filas_clases or [["", "", "", "", "", "", "", ""]]:
        clases.append(fila)
    fin_clases = 4 + max(1, len(filas_clases))
    agregar_tabla(clases, f"A4:H{fin_clases}", "DetalleClases")
    for columna, ancho in zip("ABCDEFGH", [24, 12, 20, 22, 45, 65, 28, 22]):
        clases.column_dimensions[columna].width = ancho
    clases.freeze_panes = "A5"
    for fila in clases.iter_rows(min_row=5, max_row=fin_clases, min_col=1, max_col=8):
        lineas_temas = (len(str(fila[4].value or "")) + 44) // 45
        lineas_resumen = (len(str(fila[5].value or "")) + 64) // 65
        clases.row_dimensions[fila[0].row].height = min(
            220, 16 * max(3, lineas_temas, lineas_resumen)
        )
        for celda in fila:
            celda.font = Font(name="Arial", size=10, color=TEXTO_OSCURO)
            celda.alignment = Alignment(vertical="top", wrap_text=celda.column in (3, 4, 5, 6))
        fila[7].number_format = "yyyy-mm-dd hh:mm"

    configurar_hoja(temas_hoja, "Temas por idioma enseñado")
    encabezados_temas = ["Idioma enseñado", "Tema", "Archivo", "Idioma hablado"]
    for columna, valor in enumerate(encabezados_temas, 1):
        temas_hoja.cell(row=4, column=columna, value=valor)
    estilo_encabezado(temas_hoja[4][0:len(encabezados_temas)])
    for fila in filas_temas or [["", "", "", ""]]:
        temas_hoja.append(fila)
    fin_temas = 4 + max(1, len(filas_temas))
    agregar_tabla(temas_hoja, f"A4:D{fin_temas}", "TemasIdioma")
    for columna, ancho in zip("ABCD", [22, 48, 24, 22]):
        temas_hoja.column_dimensions[columna].width = ancho
    temas_hoja.freeze_panes = "A5"
    for fila in temas_hoja.iter_rows(min_row=5, max_row=fin_temas, min_col=1, max_col=4):
        temas_hoja.row_dimensions[fila[0].row].height = 30
        for celda in fila:
            celda.font = Font(name="Arial", size=10, color=TEXTO_OSCURO)
            celda.alignment = Alignment(vertical="top", wrap_text=celda.column in (2, 4))

    temporal = ruta.with_name(f"{ruta.stem}.tmp{ruta.suffix}")
    libro.save(temporal)
    temporal.replace(ruta)


def main() -> int:
    args = leer_argumentos()
    base = Path(__file__).resolve().parent
    entrada = resolver(base, args.entrada)
    salida = resolver(base, args.salida)
    datos = salida / "_datos"
    transcripciones = datos / "transcripciones"
    detalles_dir = datos / "detalles"
    ruta_estado = datos / "estado.json"
    salida_clases = salida / "clases"
    ruta_excel = salida / "clases.xlsx"

    load_dotenv(base / ".env")
    if not os.getenv("GROQ_API_KEY"):
        print("ERROR: falta GROQ_API_KEY. Revisa el archivo .env.", file=sys.stderr)
        return 2

    entrada.mkdir(parents=True, exist_ok=True)
    transcripciones.mkdir(parents=True, exist_ok=True)
    detalles_dir.mkdir(parents=True, exist_ok=True)
    salida_clases.mkdir(parents=True, exist_ok=True)

    formatos_admitidos = FORMATOS_AUDIO | FORMATOS_TEXTO
    entradas = sorted(
        ruta for ruta in entrada.rglob("*")
        if ruta.is_file() and ruta.suffix.lower() in formatos_admitidos
    )
    if not entradas:
        crear_excel(ruta_excel, cargar_detalles_nuevos(detalles_dir))
        print(f"No encontré audios ni textos en {entrada}.")
        return 0

    cliente = Groq(api_key=os.environ["GROQ_API_KEY"])
    estado = leer_estado(ruta_estado)
    exitos = errores = 0

    for numero, archivo_entrada in enumerate(entradas, 1):
        try:
            clave = str(archivo_entrada.resolve())
            firma_actual = firma(archivo_entrada)
            hash_actual = hash_contenido(archivo_entrada)
            registro = estado["procesados"].get(clave, {})
            sin_cambios = registro.get("firma") == firma_actual

            if not args.reprocesar and sin_cambios and registro.get("duplicado_de"):
                print(f"[{numero}/{len(entradas)}] Duplicado omitido: {archivo_entrada.name}")
                continue
            if (
                not args.reprocesar
                and sin_cambios
                and registro.get("version_analisis") == VERSION_ANALISIS
            ):
                print(f"[{numero}/{len(entradas)}] Ya procesado: {archivo_entrada.name}")
                continue
            if not args.reprocesar and not sin_cambios:
                duplicado_de = next(
                    (
                        ruta_anterior
                        for ruta_anterior, datos_anteriores in estado["procesados"].items()
                        if ruta_anterior != clave
                        and datos_anteriores.get("hash_contenido") == hash_actual
                    ),
                    None,
                )
                if duplicado_de:
                    estado["procesados"][clave] = {
                        "firma": firma_actual,
                        "hash_contenido": hash_actual,
                        "duplicado_de": duplicado_de,
                        "version_analisis": VERSION_ANALISIS,
                        "omitido_utc": datetime.now(timezone.utc).isoformat(),
                    }
                    guardar_json(ruta_estado, estado)
                    print(f"[{numero}/{len(entradas)}] Duplicado omitido: {archivo_entrada.name}")
                    continue

            es_audio = archivo_entrada.suffix.lower() in FORMATOS_AUDIO
            tipo_entrada = "audio" if es_audio else "texto"
            nombre_base = f"{archivo_entrada.stem}_{identificador_archivo(archivo_entrada)}"
            ruta_texto = transcripciones / f"{nombre_base}.txt"

            if not args.reprocesar and sin_cambios and ruta_texto.exists():
                print(f"[{numero}/{len(entradas)}] Reutilizando transcripción de {archivo_entrada.name}...")
                texto = leer_texto(ruta_texto)
                idioma_sugerido = normalizar_idioma(registro.get("idioma", ""))
            elif es_audio:
                print(f"[{numero}/{len(entradas)}] Transcribiendo {archivo_entrada.name}...")
                texto, idioma_sugerido = transcribir(cliente, archivo_entrada, args.idioma)
            else:
                print(f"[{numero}/{len(entradas)}] Leyendo texto {archivo_entrada.name}...")
                texto = leer_texto(archivo_entrada)
                idioma_sugerido = "desconocido"

            print(f"[{numero}/{len(entradas)}] Analizando clase {archivo_entrada.name}...")
            analisis = analizar_clase(
                cliente,
                texto,
                idioma_sugerido,
                archivo_entrada.name,
            )
            momento = datetime.now(timezone.utc).isoformat()
            ruta_texto.write_text(texto, encoding="utf-8")
            detalle = {
                "version_analisis": VERSION_ANALISIS,
                "archivo_original": clave,
                "tipo_entrada": tipo_entrada,
                "hash_contenido": hash_actual,
                **analisis,
                "idiomas_hablados_nombres": [
                    nombre_idioma(x) for x in analisis["idiomas_hablados"]
                ],
                "idiomas_ensenados_nombres": [
                    nombre_idioma(x) for x in analisis["idiomas_ensenados"]
                ],
                "transcripcion": str(ruta_texto),
                "modelo_transcripcion": MODELO_AUDIO if es_audio else None,
                "modelo_clasificacion": MODELO_TEXTO,
                "procesado_utc": momento,
            }
            guardar_json(detalles_dir / f"{nombre_base}.json", detalle)
            if not args.sin_copiar:
                copiar_por_idioma(
                    archivo_entrada,
                    salida_clases,
                    analisis["idiomas_ensenados"],
                    hash_actual,
                )
            estado["procesados"][clave] = {
                "firma": firma_actual,
                "tipo_entrada": tipo_entrada,
                "hash_contenido": hash_actual,
                "version_analisis": VERSION_ANALISIS,
                "idiomas_hablados": analisis["idiomas_hablados"],
                "idiomas_ensenados": analisis["idiomas_ensenados"],
                "procesado_utc": momento,
            }
            guardar_json(ruta_estado, estado)
            exitos += 1
            idiomas = ", ".join(nombre_idioma(x) for x in analisis["idiomas_ensenados"])
            print(
                f"[{numero}/{len(entradas)}] Listo: {tipo_entrada} / "
                f"enseña {idiomas} / {len(analisis['temas'])} tema(s)"
            )
        except Exception as error:
            errores += 1
            print(f"[{numero}/{len(entradas)}] ERROR: {error}", file=sys.stderr)

    try:
        crear_excel(ruta_excel, cargar_detalles_nuevos(detalles_dir))
        print(f"Excel actualizado: {ruta_excel}")
    except PermissionError:
        errores += 1
        print("ERROR: cierra clases.xlsx en Excel y vuelve a ejecutar.", file=sys.stderr)
    except Exception as error:
        errores += 1
        print(f"ERROR al crear clases.xlsx: {error}", file=sys.stderr)

    print(f"Terminado. Actualizados: {exitos}. Errores: {errores}. Salida: {salida}")
    return 1 if errores else 0


if __name__ == "__main__":
    raise SystemExit(main())
